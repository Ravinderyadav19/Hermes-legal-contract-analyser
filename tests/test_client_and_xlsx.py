import pytest

from hermes_legal.analysis.engine import analyze_contract
from hermes_legal.memory.store import MemoryStore
from hermes_legal.providers.offline_provider import OfflineProvider

TEXT_A = "FREELANCE SERVICE AGREEMENT between TechCorp and Contractor. 1. Liability. Uncapped and unlimited liability."
TEXT_B = "MUTUAL NON-DISCLOSURE AGREEMENT between Acme and Beta LLC. 1. Confidentiality survives perpetually."


def test_client_tag_stored_and_filterable(tmp_path):
    memory = MemoryStore(base_dir=tmp_path)
    analyze_contract(TEXT_A, provider=OfflineProvider(), memory=memory, client="Acme Corp")
    analyze_contract(TEXT_B, provider=OfflineProvider(), memory=memory, client="Beta LLC")

    assert set(memory.clients()) == {"Acme Corp", "Beta LLC"}
    assert len(memory.find_by_client("Acme Corp")) == 1
    assert len(memory.find_by_client("Beta LLC")) == 1
    assert len(memory.find_by_client("Nonexistent")) == 0


def test_memory_store_respects_env_var(tmp_path, monkeypatch=None):
    import os
    old = os.environ.get("HERMES_LEGAL_HOME")
    os.environ["HERMES_LEGAL_HOME"] = str(tmp_path)
    try:
        memory = MemoryStore()
        assert memory.base_dir == tmp_path
    finally:
        if old is None:
            os.environ.pop("HERMES_LEGAL_HOME", None)
        else:
            os.environ["HERMES_LEGAL_HOME"] = old


def test_write_batch_xlsx(tmp_path):
    pytest.importorskip("openpyxl")
    from hermes_legal.reports.xlsx_export import write_batch_xlsx

    rows = [
        {"file": "a.txt", "contract_type": "NDA", "overall_risk": "LOW", "verdict": "SIGN",
         "red_flag_count": 0, "clause_count": 5, "provider": "offline", "hash": "abc"},
        {"file": "b.txt", "contract_type": "Freelance", "overall_risk": "CRITICAL", "verdict": "REJECT",
         "red_flag_count": 3, "clause_count": 6, "provider": "offline", "hash": "def"},
    ]
    out_path = tmp_path / "batch.xlsx"
    saved = write_batch_xlsx(rows, out_path)
    assert saved is not None and saved.exists()

    from openpyxl import load_workbook
    wb = load_workbook(str(saved))
    ws = wb.active
    assert ws.max_row == 3
    assert ws.cell(row=1, column=1).value == "File"
