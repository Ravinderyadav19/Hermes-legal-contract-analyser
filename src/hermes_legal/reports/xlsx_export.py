from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

RISK_FILL = {"CRITICAL": "F8CBCB", "HIGH": "FCE4C4", "MEDIUM": "FCF3C4", "LOW": "D4F1DC"}
VERDICT_FILL = {"SIGN": "D4F1DC", "NEGOTIATE": "FCF3C4", "REJECT": "F8CBCB"}


def write_batch_xlsx(rows: List[Dict[str, Any]], out_path: str | Path) -> Optional[Path]:
    """
    Write a formatted, color-coded Excel workbook summarizing a batch run
    - the format most law firms and clients actually work in day to day,
    as opposed to CSV. Returns None (without raising) if openpyxl is not
    installed, keeping this a purely optional feature.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    out_path = Path(out_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Batch Summary"

    headers = ["File", "Contract Type", "Parties", "Language", "Overall Risk", "Verdict",
               "Red Flags", "Clauses", "Provider", "Hash"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50")
        cell.alignment = Alignment(horizontal="center")

    key_order = ["file", "contract_type", "parties", "language", "overall_risk", "verdict",
                 "red_flag_count", "clause_count", "provider", "hash"]
    for row in rows:
        ws.append([row.get(k, "") for k in key_order])
        r = ws.max_row
        risk = row.get("overall_risk", "")
        verdict = row.get("verdict", "")
        if risk in RISK_FILL:
            ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=RISK_FILL[risk])
        if verdict in VERDICT_FILL:
            ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor=VERDICT_FILL[verdict])

    for i, header in enumerate(headers, start=1):
        max_len = max([len(header)] + [len(str(r.get(key_order[i - 1], ""))) for r in rows]) if rows else len(header)
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 10), 45)

    ws.freeze_panes = "A2"
    wb.save(str(out_path))
    return out_path
