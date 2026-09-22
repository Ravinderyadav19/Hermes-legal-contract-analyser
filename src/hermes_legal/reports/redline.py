from __future__ import annotations

from pathlib import Path
from typing import Optional

from ..providers import AnalysisResult


def render_redline_markdown(result: AnalysisResult) -> str:
    """A focused, negotiation-ready view: only the clauses that need changes."""
    flagged = [c for c in result.clauses if c.get("is_red_flag") or c.get("negotiation_suggestion")]
    lines = [f"# Redline Suggestions - {result.contract_type}", ""]
    if not flagged:
        lines.append("No clauses were flagged for renegotiation.")
        return "\n".join(lines)

    for c in flagged:
        lines += [
            f"## {c.get('name')}  (risk {c.get('score')}/10)",
            "",
            f"**Issue:** {c.get('finding', '')}",
            "",
            f"**Suggested replacement language:**",
            f"> {c.get('negotiation_suggestion') or 'Consult an attorney for specific replacement language.'}",
            "",
        ]
    return "\n".join(lines)


def write_redline_docx(result: AnalysisResult, out_path: str | Path) -> Optional[Path]:
    """
    Write a Word document with one heading + issue + suggestion per
    flagged clause, styled so it reads like a redline memo you can attach
    to an email back to the counter-party. Returns None (and does not
    raise) if python-docx is not installed, so this stays a purely
    optional feature.
    """
    try:
        import docx
        from docx.shared import Pt, RGBColor
    except ImportError:
        return None

    out_path = Path(out_path)
    document = docx.Document()
    document.add_heading(f"Redline Suggestions - {result.contract_type}", level=1)
    document.add_paragraph(f"Overall risk: {result.overall_risk}  |  Verdict: {result.verdict}")

    flagged = [c for c in result.clauses if c.get("is_red_flag") or c.get("negotiation_suggestion")]
    if not flagged:
        document.add_paragraph("No clauses were flagged for renegotiation.")
    for c in flagged:
        document.add_heading(f"{c.get('name')} (risk {c.get('score')}/10)", level=2)
        p = document.add_paragraph()
        run = p.add_run("Issue: ")
        run.bold = True
        p.add_run(c.get("finding", ""))

        p2 = document.add_paragraph()
        run2 = p2.add_run("Suggested replacement language: ")
        run2.bold = True
        run3 = p2.add_run(c.get("negotiation_suggestion") or "Consult an attorney for specific language.")
        run3.italic = True
        run3.font.color.rgb = RGBColor(0x1A, 0x73, 0x2E)

    document.add_paragraph()
    footer = document.add_paragraph(
        "Hermes Legal Advisor provides contract analysis, not legal advice. "
        "Always consult a qualified attorney before signing any contract."
    )
    footer.runs[0].font.size = Pt(8)

    document.save(str(out_path))
    return out_path


def write_redline_docx_inline(original_path: str | Path, result: AnalysisResult, out_path: str | Path) -> Optional[Path]:
    """
    Take the user's actual original .docx file and insert Hermes'
    suggestions as real Word "Tracked Changes" insertions, directly after
    the paragraph that appears to match each flagged clause - not a
    separate memo, and not just colored text. The output opens in Word's
    Review pane like any other tracked edit: it can be accepted, rejected,
    or commented on paragraph by paragraph.

    Matching is a simple case-insensitive keyword search for the clause
    name, so it works without needing exact paragraph boundaries from the
    provider. Falls back gracefully (returns None) if python-docx isn't
    installed or no clause could be located in the document text.
    """
    try:
        import docx
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError:
        return None

    import datetime

    original_path = Path(original_path)
    out_path = Path(out_path)
    document = docx.Document(str(original_path))
    now_iso = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    def _tracked_insertion_paragraph(text: str, change_id: int):
        p = OxmlElement("w:p")
        ins = OxmlElement("w:ins")
        ins.set(qn("w:id"), str(change_id))
        ins.set(qn("w:author"), "Hermes Legal Advisor")
        ins.set(qn("w:date"), now_iso)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        italic = OxmlElement("w:i")
        rpr.append(italic)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        ins.append(r)
        p.append(ins)
        return p

    flagged = [c for c in result.clauses if c.get("is_red_flag") or c.get("negotiation_suggestion")]
    inserted = 0
    change_id = 1000
    for c in flagged:
        name = c.get("name", "")
        keywords = [name.lower()] + name.lower().replace("-", " ").split()
        target = None
        for p in document.paragraphs:
            text_lower = p.text.lower()
            if any(kw in text_lower for kw in keywords if len(kw) > 3):
                target = p
                break
        if target is None:
            continue

        note_text = (
            f"[Hermes Legal Advisor - {name}, risk {c.get('score')}/10] {c.get('finding', '')} "
            f"Suggested: {c.get('negotiation_suggestion') or 'consult an attorney for specific language.'}"
        )
        new_p = _tracked_insertion_paragraph(note_text, change_id)
        change_id += 1
        target._p.addnext(new_p)
        inserted += 1

    if inserted == 0:
        return None

    document.save(str(out_path))
    return out_path
