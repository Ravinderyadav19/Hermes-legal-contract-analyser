"""
Document readers for Hermes Legal Advisor.

Real contracts rarely arrive as .txt files. This module adds first-class
support for PDF and DOCX so the tool is actually usable on documents
people receive by email, in addition to the plain text files used in the
sample/demo set.
"""

from __future__ import annotations

from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class UnsupportedFileError(ValueError):
    pass


def read_document(path: str | Path) -> str:
    """Read a contract file and return its plain text content."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {p}")

    suffix = p.suffix.lower()
    if suffix in (".txt", ".md"):
        return _read_txt(p)
    if suffix == ".pdf":
        return _read_pdf(p)
    if suffix == ".docx":
        return _read_docx(p)

    raise UnsupportedFileError(
        f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _read_txt(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="replace")


def _read_pdf(p: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError(
            "Reading PDF files requires pypdf. Install with: pip install pypdf"
        ) from exc

    reader = PdfReader(str(p))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    combined = "\n\n".join(pages).strip()
    if not combined:
        raise ValueError(
            f"No extractable text found in {p.name}. It may be a scanned image PDF "
            "that needs OCR before it can be analyzed."
        )
    return combined


def _read_docx(p: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            "Reading DOCX files requires python-docx. Install with: pip install python-docx"
        ) from exc

    document = docx.Document(str(p))
    parts = [para.text for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
